"""
dashboard_server.py — Job Hunter Agent Dashboard Server
Run: python dashboard_server.py
Open: http://localhost:8080
"""

import json
import signal
import os
import csv
import io
import re
from http.server import HTTPServer, BaseHTTPRequestHandler
from urllib.parse import urlparse
from career_tracker import load_career_apps, load_career_sites, save_career_apps, save_career_sites

# ---------------------------------------------------------------------------
# Import all modules
# ---------------------------------------------------------------------------
from dashboard_ui import html_page, render_pagination_js
from dashboard_profile import build_profile_html, build_profile_js
from dashboard_career import build_career_html, build_career_js
from dashboard_jobodyssey import build_jobodyssey_html, build_jobodyssey_js
from dashboard_scraper import build_scraper_html, build_scraper_js
from dashboard_log import log_event, add_tab_log, get_tab_logs
from jobodyssey_store import load_jobs, count_jobs, get_job, get_stats, get_filter_options, update_status, get_description_locations
import workflow_store as ws
import workflow_engine as we
from dashboard_workflow_detail import build_workflow_detail_page

# ---------------------------------------------------------------------------
# Profile data (profile.json) — moved here from the removed
# dashboard_workflows.py, which used to own all shared file paths.
# ---------------------------------------------------------------------------
BASE_DIR     = os.path.dirname(os.path.abspath(__file__))
PROFILE_JSON = os.path.join(BASE_DIR, 'profile.json')

def read_profile():
    if not os.path.exists(PROFILE_JSON):
        return {}
    try:
        with open(PROFILE_JSON, 'r') as f:
            return json.load(f)
    except Exception:
        return {}

# ---------------------------------------------------------------------------
# XLSX export — used by /api/export-file. Color-coding is picked by which
# recognizable column is present: 'status' (Job Odyssey) or 'stage'
# (Applications). Job Sites and any export without either column just gets
# the bold frozen header with no row coloring.
# ---------------------------------------------------------------------------
JO_STATUS_FILL = {
    'eligible':   'C8E6C9',  # light green
    'applied':    'BBDEFB',  # light blue
    'hold':       'FFF9C4',  # light yellow
    'ineligible': 'FFCDD2',  # light red
    'blacklist':  'E0E0E0',  # gray
}
APP_STAGE_FILL = {
    'Offer':       'C8E6C9',
    'Interview':   'BBDEFB',
    'In Progress': 'E1F5FE',
    'Applied':     'FFF9C4',
    'Rejected':    'FFCDD2',
    'Ghosted':     'E0E0E0',
    'Withdrawn':   'ECEFF1',
    'Closed':      'ECEFF1',
}

def _build_xlsx_bytes(rows, keys, labels):
    # write_only=True streams the file instead of building the whole
    # worksheet as in-memory Cell objects — openpyxl's own recommendation
    # for anything beyond a few hundred rows. The normal (non-write-only)
    # mode was measured taking ~4s for ~2,900 rows x 12 cols due to
    # per-cell object overhead; this mode does the same in well under 1s.
    from openpyxl import Workbook
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    wb = Workbook(write_only=True)
    ws = wb.create_sheet('Export')
    # In write_only mode, freeze_panes must be set BEFORE the first append()
    # call or it silently fails to persist on save — confirmed by testing.
    ws.freeze_panes = 'A2'

    header_font = Font(name='Arial', bold=True, color='FFFFFF')
    header_fill = PatternFill('solid', fgColor='3949AB')
    header_align = Alignment(vertical='center')
    body_font = Font(name='Arial')

    header_cells = []
    for label in labels:
        c = WriteOnlyCell(ws, value=label)
        c.font = header_font
        c.fill = header_fill
        c.alignment = header_align
        header_cells.append(c)
    ws.append(header_cells)

    status_idx = keys.index('status') if 'status' in keys else None
    stage_idx  = keys.index('stage') if 'stage' in keys else None
    fill_cache = {}  # hex -> PatternFill, reused across rows sharing a color

    for row in rows:
        fill_hex = None
        if status_idx is not None:
            fill_hex = JO_STATUS_FILL.get(str(row.get('status', '')).lower())
        elif stage_idx is not None:
            fill_hex = APP_STAGE_FILL.get(str(row.get('stage', '')))
        fill = None
        if fill_hex:
            fill = fill_cache.get(fill_hex)
            if fill is None:
                fill = PatternFill('solid', fgColor=fill_hex)
                fill_cache[fill_hex] = fill

        cells = []
        for k in keys:
            c = WriteOnlyCell(ws, value=row.get(k, ''))
            c.font = body_font
            if fill:
                c.fill = fill
            cells.append(c)
        ws.append(cells)

    for i, label in enumerate(labels, start=1):
        ws.column_dimensions[get_column_letter(i)].width = max(12, min(40, len(str(label)) + 6))

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# Main dashboard HTML builder
# ---------------------------------------------------------------------------
def build_dashboard_html(active_tab='career'):
    profile   = read_profile()
    prof_json = json.dumps(profile)

    body = '''
<div class="header">
  <h1>🌌 Bala's Solaris</h1>
  <div class="nav">
    <button class="''' + ('active' if active_tab=='jobodyssey' else '') + '''" onclick="showPage(\'jobodyssey\',this)">🪐 Job Odyssey</button>
    <button class="''' + ('active' if active_tab=='career' else '') + '''" onclick="showPage(\'career\',this)">🚀 Career Voyager</button>
    <button class="''' + ('active' if active_tab=='profile' else '') + '''" onclick="showPage(\'profile\',this)">⭐ Profile Nexus</button>
    <button class="''' + ('active' if active_tab=='scraper' else '') + '''" onclick="showPage(\'scraper\',this)">🕳️ Scraper Abyss</button>
  </div>
</div>
<div class="status-bar" id="status_bar">Ready.</div>
''' + build_jobodyssey_html() + '''
''' + build_career_html() + '''
''' + build_profile_html(prof_json) + '''
''' + build_scraper_html() + '''
<script>
''' + render_pagination_js() + '''
''' + build_jobodyssey_js() + '''
''' + build_career_js() + '''
''' + build_profile_js(prof_json) + '''
''' + build_scraper_js() + '''
// ── Page nav ───────────────────────────────────────────────────────────────────
var TAB_TITLES = {jobodyssey: 'Job Odyssey', career: 'Career Voyager', profile: 'Profile Nexus', scraper: 'Scraper Abyss'};
function showPage(name,btn,updateUrl){
  if(updateUrl===undefined) updateUrl = true;
  document.querySelectorAll('.page').forEach(p=>p.classList.remove('active'));
  document.querySelectorAll('.nav button').forEach(b=>b.classList.remove('active'));
  document.getElementById('page-'+name).classList.add('active');
  if(!btn) btn = [...document.querySelectorAll('.nav button')].find(b=>b.getAttribute('onclick')&&b.getAttribute('onclick').includes("'"+name+"'"));
  if(btn)btn.classList.add('active');
  document.title = "Bala's Solaris — " + (TAB_TITLES[name] || name);
  if(updateUrl && window.location.pathname !== '/'+name){
    history.pushState({tab:name}, '', '/'+name);
  }
  if(name==='career') crLoad();
  if(name==='jobodyssey') joLoad();
}
window.addEventListener('popstate', function(e){
  const name = (e.state && e.state.tab) || 'jobodyssey';
  showPage(name, null, false);
});
window.onload = function() {
  fillProfileForm(PROFILE_DATA);
  // Activate tab from URL path
  const path = window.location.pathname;
  let activeTab = 'jobodyssey';
  if(path.startsWith('/career')) activeTab = 'career';
  else if(path.startsWith('/profile')) activeTab = 'profile';
  else if(path.startsWith('/scraper')) activeTab = 'scraper';
  history.replaceState({tab:activeTab}, '', '/'+activeTab);
  showPage(activeTab, null, false);
};
</script>
'''
    tab_titles = {
        'jobodyssey': 'Job Odyssey',
        'career': 'Career Voyager',
        'profile': 'Profile Nexus',
        'scraper': 'Scraper Abyss',
    }
    page_title = "Bala's Solaris — " + tab_titles.get(active_tab, active_tab)
    return html_page(page_title, body)

# ---------------------------------------------------------------------------
# HTTP Handler
# ---------------------------------------------------------------------------
class DashboardHandler(BaseHTTPRequestHandler):

    def log_message(self, format, *args):
        if args and str(args[1]) not in ('200', '304'):
            print('[Server] ' + format % args, flush=True)

    def send_json(self, data, code=200):
        try:
            body = json.dumps(data).encode()
            self.send_response(code)
            self.send_header('Content-Type', 'application/json')
            self.send_header('Content-Length', len(body))
            self.end_headers()
            self.wfile.write(body)
        except (ConnectionAbortedError, BrokenPipeError, ConnectionResetError):
            pass  # Browser closed connection — ignore silently

    def send_html(self, html, code=200):
        try:
            body = html.encode('utf-8')
            self.send_response(code)
            self.send_header('Content-Type', 'text/html; charset=utf-8')
            self.send_header('Content-Length', len(body))
            self.end_headers()
            self.wfile.write(body)
        except (ConnectionAbortedError, BrokenPipeError, ConnectionResetError):
            pass

    def send_file_download(self, data_bytes, filename, content_type='text/csv'):
        try:
            self.send_response(200)
            self.send_header('Content-Type', content_type + '; charset=utf-8')
            self.send_header('Content-Disposition', 'attachment; filename="' + filename + '"')
            self.send_header('Content-Length', len(data_bytes))
            self.end_headers()
            self.wfile.write(data_bytes)
        except (ConnectionAbortedError, BrokenPipeError, ConnectionResetError):
            pass

    def read_body(self):
        try:
            length = int(self.headers.get('Content-Length', 0))
            if not length:
                return {}
            raw = self.rfile.read(length)
            return json.loads(raw)
        except (json.JSONDecodeError, ValueError):
            self.send_json({'ok': False, 'error': 'Browser is busy, please try again in a moment.'}, 503)
            return None
        except Exception:
            return {}

    def do_GET(self):
        parsed = urlparse(self.path)
        path   = parsed.path

        if path in ('/', '/dashboard'):
            self.send_html(build_dashboard_html('jobodyssey'))

        elif path == '/jobodyssey':
            self.send_html(build_dashboard_html('jobodyssey'))

        elif path == '/career':
            self.send_html(build_dashboard_html('career'))

        elif path == '/profile':
            self.send_html(build_dashboard_html('profile'))

        elif path == '/scraper':
            self.send_html(build_dashboard_html('scraper'))

        elif path == '/workflow-detail':
            from urllib.parse import parse_qs
            qs = parse_qs(parsed.query)
            wf_id = qs.get('id', [''])[0]
            if not ws.load_workflow(wf_id):
                self.send_html('<h1>404 — workflow not found</h1>', 404)
            else:
                self.send_html(build_workflow_detail_page(wf_id))

        elif path == '/api/career-load':
            self.send_json({
                'apps':  load_career_apps(),
                'sites': load_career_sites()
            })

        elif path == '/api/workflow-list':
            self.send_json({'workflows': ws.list_workflows()})

        elif path == '/api/workflow-detail':
            from urllib.parse import parse_qs
            qs = parse_qs(parsed.query)
            wf_id = qs.get('id', [''])[0]
            data = ws.load_workflow(wf_id)
            if data is None:
                self.send_json({'ok': False, 'error': 'Workflow not found'}, 404)
            else:
                self.send_json({'ok': True, 'workflow': data})

        elif path == '/api/workflow-logs':
            from urllib.parse import parse_qs
            qs = parse_qs(parsed.query)
            wf_id  = qs.get('id', [''])[0]
            search = qs.get('search', [''])[0]
            self.send_json({'logs': ws.wf_get_logs(wf_id, search=search or None)})

        elif path == '/api/jobodyssey-load':
            from urllib.parse import parse_qs
            qs = parse_qs(parsed.query)
            status_raw  = qs.get('status', [''])[0]
            status_list = [s for s in status_raw.split(',') if s] if status_raw else []
            status_op   = qs.get('statusOp', ['or'])[0]
            if status_op not in ('or', 'not'):
                status_op = 'or'
            portal      = qs.get('portal', [''])[0]
            location    = qs.get('location', [''])[0]
            search      = qs.get('search', [''])[0]
            sort        = qs.get('sort', ['date_desc'])[0]
            validation  = qs.get('validation', [''])[0]
            link_type   = qs.get('linkType', [''])[0]
            try:
                page      = max(1, int(qs.get('page', ['1'])[0]))
                page_size = max(1, min(200, int(qs.get('pageSize', ['20'])[0])))
            except ValueError:
                page, page_size = 1, 20
            filt = dict(status=status_list or None, status_op=status_op, portal=portal or None, location=location or None,
                        search=search or None, validation=validation or None, link_type=link_type or None)
            self.send_json({
                'jobs': load_jobs(sort=sort, limit=page_size, offset=(page - 1) * page_size, **filt),
                'total': count_jobs(**filt),
                'stats': get_stats(portal=portal or None, location=location or None, search=search or None,
                                    validation=validation or None, link_type=link_type or None),
                'filterOptions': get_filter_options()
            })

        elif path == '/api/jobodyssey-export':
            # Same filters as jobodyssey-load, but NO page/pageSize — this
            # returns every matching row for the Download-export feature,
            # since exports must never be capped to a single visible page.
            from urllib.parse import parse_qs
            qs = parse_qs(parsed.query)
            status_raw  = qs.get('status', [''])[0]
            status_list = [s for s in status_raw.split(',') if s] if status_raw else []
            status_op   = qs.get('statusOp', ['or'])[0]
            if status_op not in ('or', 'not'):
                status_op = 'or'
            portal      = qs.get('portal', [''])[0]
            location    = qs.get('location', [''])[0]
            search      = qs.get('search', [''])[0]
            sort        = qs.get('sort', ['date_desc'])[0]
            validation  = qs.get('validation', [''])[0]
            link_type   = qs.get('linkType', [''])[0]
            filt = dict(status=status_list or None, status_op=status_op, portal=portal or None, location=location or None,
                        search=search or None, validation=validation or None, link_type=link_type or None)
            self.send_json({'jobs': load_jobs(sort=sort, **filt)})

        elif path == '/api/jobodyssey-job-detail':
            from urllib.parse import parse_qs
            qs = parse_qs(parsed.query)
            try:
                job_id = int(qs.get('id', [''])[0])
            except (ValueError, TypeError):
                self.send_json({'ok': False, 'error': 'Invalid id'}, 400)
            else:
                job = get_job(job_id)
                if job is None:
                    self.send_json({'ok': False, 'error': 'Job not found'}, 404)
                else:
                    self.send_json({'ok': True, 'job': job})

        elif path == '/api/logs':
            from urllib.parse import parse_qs
            qs = parse_qs(parsed.query)
            tab = qs.get('tab', ['career'])[0]
            self.send_json({'logs': get_tab_logs(tab)})

        else:
            self.send_html('<h1>404</h1>', 404)

    def do_POST(self):
        path = urlparse(self.path).path
        body = self.read_body()
        if body is None:
            return  # read_body already sent 503

        if path == '/api/save-profile':
            try:
                with open(PROFILE_JSON, 'w') as f:
                    json.dump(body, f, indent=4)
                self.send_json({'ok': True})
            except Exception as e:
                log_event('Server', 'Save profile error: ' + str(e)[:80])
                self.send_json({'ok': False, 'error': str(e)})

        elif path == '/api/career-save':
            try:
                ok_a = save_career_apps(body.get('apps', []))
                ok_s = save_career_sites(body.get('sites', []))
                if not (ok_a and ok_s):
                    log_event('Server', 'Career save failed (see CareerTracker log above)')
                self.send_json({'ok': ok_a and ok_s})
            except Exception as e:
                log_event('Server', 'Career save error: ' + str(e)[:80])
                self.send_json({'ok': False, 'error': str(e)[:50]})

        elif path == '/api/scraper-submit':
            try:
                engines   = body.get('engines', [])
                portals   = body.get('portals', [])
                keywords  = body.get('keywords', [])
                locations = body.get('locations', [])
                blacklist = body.get('blacklist', [])
                created_ids = []
                prev_wf_id = None    # each workflow's depends_on target
                first_started = False  # only the very first workflow in the
                                        # whole chain is started here; every
                                        # other one sits in 'waiting' until
                                        # workflow_engine's _trigger_dependents
                                        # wakes it after its dependency finishes

                for portal in portals:
                    scraper_wf_id = None

                    if 'Scraper' in engines:
                        scraper_wf_id = ws.create_workflow(portal, keywords, locations,
                                                            blacklist=blacklist, depends_on=prev_wf_id)
                        created_ids.append(scraper_wf_id)
                        if not first_started:
                            we.start_workflow(scraper_wf_id)
                            first_started = True
                        prev_wf_id = scraper_wf_id

                    # JSearch and JobsPipe both return complete descriptions
                    # from Scraper directly — no separate detail page to
                    # visit, so Descriptor is always skipped for them.
                    NO_DESCRIPTOR_PORTALS = ('JSearch', 'JobsPipe')
                    if 'Descriptor' in engines and portal not in NO_DESCRIPTOR_PORTALS:
                        if locations:
                            desc_wf_id = ws.create_descriptor_workflow(
                                portal, locations, depends_on=prev_wf_id
                            )
                            created_ids.append(desc_wf_id)
                            if not first_started:
                                we.start_workflow(desc_wf_id)
                                first_started = True
                            prev_wf_id = desc_wf_id
                        else:
                            log_event('Server', 'Descriptor skipped for ' + portal + ' — no locations selected')
                    elif 'Descriptor' in engines and portal == 'JSearch':
                        log_event('Server', 'Descriptor skipped for JSearch — full description already included by Scraper')
                    elif 'Descriptor' in engines and portal == 'JobsPipe':
                        log_event('Server', 'Descriptor skipped for JobsPipe — full description already included by Scraper')

                log_event('Server', 'Scraper request created workflow(s): ' + str(created_ids))
                self.send_json({'ok': True, 'workflow_ids': created_ids})
            except Exception as e:
                log_event('Server', 'Scraper submit error: ' + str(e)[:100])
                self.send_json({'ok': False, 'error': str(e)[:80]})

        elif path == '/api/workflow-action':
            try:
                wf_id  = str(body.get('workflow_id', ''))
                action = str(body.get('action', ''))
                result = False
                if action == 'hold':
                    we.hold_workflow(wf_id); result = True
                elif action == 'cancel':
                    we.cancel_workflow(wf_id); result = True
                elif action == 'revoke':
                    we.revoke_workflow(wf_id); result = True
                elif action == 'retry':
                    result = we.retry_workflow(wf_id)
                elif action == 'reset':
                    result = we.reset_workflow(wf_id)
                self.send_json({'ok': result})
            except Exception as e:
                log_event('Server', 'Workflow action error: ' + str(e)[:100])
                self.send_json({'ok': False, 'error': str(e)[:80]})

        elif path == '/api/jobodyssey-status':
            try:
                job_id = int(body.get('id'))
                status = str(body.get('status', ''))
                ok = update_status(job_id, status)
                self.send_json({'ok': ok})
            except Exception as e:
                log_event('Server', 'JobOdyssey status error: ' + str(e)[:80])
                self.send_json({'ok': False, 'error': str(e)[:50]})

        elif path == '/api/client-log':
            # Browser-side action log for one tab ('career' or 'profile').
            # Stored only — deliberately not printed to console, since the
            # console is reserved for server-side system events.
            tab     = str(body.get('tab', 'career'))[:20]
            source  = str(body.get('source', 'Client'))[:40]
            message = str(body.get('message', ''))[:300]
            if message:
                add_tab_log(tab, source, message)
            self.send_json({'ok': True})

        elif path == '/api/export-file':
            # Generic export — deliberately knows little about which tab the
            # rows came from beyond the presence of a 'status' or 'stage'
            # column (used only to pick a color-coding scheme for XLSX).
            # columns: [{key, label}, ...] — label drives the header text,
            # key looks up each row's value. CSV via Python's csv module
            # (handles embedded commas/newlines); XLSX via openpyxl with a
            # bold frozen header and color-coded rows, built once CSV was
            # confirmed working, per the original plan.
            try:
                rows        = body.get('rows', [])
                columns_in  = body.get('columns', [])
                fmt         = str(body.get('format', 'csv')).lower()
                if fmt not in ('csv', 'xlsx'):
                    fmt = 'csv'
                filename = re.sub(r'[^A-Za-z0-9_-]', '_', str(body.get('filename', 'export'))[:80]) or 'export'

                # Accept either [{key,label}, ...] or a plain list of key
                # strings (falls back to using the key itself as the label).
                keys, labels = [], []
                for c in columns_in:
                    if isinstance(c, dict):
                        keys.append(c.get('key', ''))
                        labels.append(c.get('label', c.get('key', '')))
                    else:
                        keys.append(c)
                        labels.append(c)

                if not keys:
                    self.send_json({'ok': False, 'error': 'No columns selected'}, 400)
                elif fmt == 'xlsx':
                    data = _build_xlsx_bytes(rows, keys, labels)
                    self.send_file_download(
                        data, filename + '.xlsx',
                        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
                    )
                else:
                    buf = io.StringIO()
                    writer = csv.writer(buf)
                    writer.writerow(labels)
                    for row in rows:
                        writer.writerow([row.get(k, '') for k in keys])
                    # utf-8-sig BOM so Excel opens the file with correct encoding
                    data = buf.getvalue().encode('utf-8-sig')
                    self.send_file_download(data, filename + '.csv')
            except Exception as e:
                log_event('Server', 'Export error: ' + str(e)[:80])
                self.send_json({'ok': False, 'error': str(e)[:80]}, 500)

        else:
            self.send_json({'ok': False, 'error': 'Unknown endpoint'}, 404)

# ---------------------------------------------------------------------------
# Entry point
# ---------------------------------------------------------------------------
if __name__ == '__main__':
    port   = 8080
    server = HTTPServer(('localhost', port), DashboardHandler)

    print('=' * 45, flush=True)
    print('  Job Hunter Dashboard', flush=True)
    print('  http://localhost:' + str(port), flush=True)
    print('  Press Ctrl+C to stop', flush=True)
    print('=' * 45, flush=True)
    log_event('Server', 'Dashboard started on port ' + str(port))

    def handle_sigint(signum, frame):
        raise KeyboardInterrupt

    signal.signal(signal.SIGINT, handle_sigint)

    try:
        server.serve_forever()
    except KeyboardInterrupt:
        print('\n[Server] Stopped.', flush=True)
    finally:
        server.server_close()
